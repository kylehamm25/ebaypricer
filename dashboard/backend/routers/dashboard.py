import os
from datetime import date, datetime, timedelta, timezone
from uuid import UUID

from fastapi import APIRouter, Depends, Query
from openpyxl import load_workbook

from dashboard.backend.auth import get_current_user_id
from dashboard.backend.config import EXCEL_PATH
from dashboard.backend.database import get_db

router = APIRouter(prefix="/api/v1/dashboard", tags=["dashboard"])


def _table_exists(db, name: str) -> bool:
    return bool(
        db.execute(
            "SELECT to_regclass(%s) IS NOT NULL AS exists", (f"public.{name}",)
        ).fetchone()["exists"]
    )


def _available_months_from_excel() -> list[str] | None:
    """Read the Sold Orders sheet's Sale Date column directly, so the arrows' navigable
    range reflects the workbook (the bookkeeping source of truth) even if the backend's
    Postgres sync hasn't run since the workbook last changed. Returns None (caller falls
    back to Postgres) if the file/sheet isn't reachable."""
    if not os.path.exists(EXCEL_PATH):
        return None
    try:
        wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
        try:
            if "Sold Orders" not in wb.sheetnames:
                return None
            ws = wb["Sold Orders"]
            headers = [c.value.strip() if isinstance(c.value, str) else c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            if "Sale Date" not in headers:
                return None
            date_idx = headers.index("Sale Date")
            months: set[str] = set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                v = row[date_idx] if date_idx < len(row) else None
                if v is None:
                    continue
                if isinstance(v, (datetime, date)):
                    months.add(f"{v.year:04d}-{v.month:02d}")
                else:
                    s = str(v).strip()
                    if len(s) >= 7 and s[4] == "-":
                        months.add(s[:7])
            return sorted(months, reverse=True)
        finally:
            wb.close()
    except Exception:
        return None


def _parse_month(month: str | None, today: date) -> date:
    """YYYY-MM -> first-of-month date; falls back to the current month if missing/invalid."""
    if month:
        try:
            year, mon = (int(p) for p in month.split("-", 1))
            return date(year, mon, 1)
        except (ValueError, TypeError):
            pass
    return today.replace(day=1)


def _month_end(month_start: date) -> date:
    next_month = date(month_start.year + 1, 1, 1) if month_start.month == 12 else date(month_start.year, month_start.month + 1, 1)
    return next_month - timedelta(days=1)


@router.get("/kpis")
def get_dashboard_kpis(
    user_id: UUID = Depends(get_current_user_id),
    month: str | None = Query(None, description="YYYY-MM, defaults to the current month"),
):
    today = datetime.now(timezone.utc).date()
    month_start = _parse_month(month, today)
    month_end = _month_end(month_start)
    month_str = month_start.strftime("%Y-%m")
    is_current_month = month_start.year == today.year and month_start.month == today.month
    trend_end = today if is_current_month else month_end

    with get_db() as db:
        sold_exists = _table_exists(db, "sold_orders")
        active_exists = _table_exists(db, "active_listings")

        sold_items = 0
        revenue = 0
        shipping = 0
        fees = 0
        trends = []
        top_items = []
        available_months: list[str] = []

        if sold_exists:
            available_months = _available_months_from_excel()
            if available_months is None:
                available_months = [
                    r["m"]
                    for r in db.execute(
                        """SELECT DISTINCT to_char(sale_date, 'YYYY-MM') AS m
                           FROM sold_orders WHERE user_id = %s AND sale_date IS NOT NULL
                           ORDER BY m DESC""",
                        [user_id],
                    ).fetchall()
                ]
            if month_str not in available_months:
                available_months = sorted({*available_months, month_str}, reverse=True)

            rows = db.execute(
                """SELECT sale_date AS date, COUNT(*) AS count,
                          COALESCE(SUM(item_price), 0) AS revenue
                   FROM sold_orders
                   WHERE user_id = %s AND sale_date >= %s AND sale_date <= %s
                   GROUP BY sale_date ORDER BY sale_date ASC""",
                [user_id, month_start, trend_end],
            ).fetchall()
            by_date = {r["date"].isoformat(): r for r in rows}
            trends = []
            d = month_start
            while d <= trend_end:
                ds = d.isoformat()
                r = by_date.get(ds)
                trends.append(
                    {
                        "date": ds,
                        "count": r["count"] if r else 0,
                        "revenue": r["revenue"] if r else 0,
                    }
                )
                d += timedelta(days=1)

            month_filter = "to_char(sale_date, 'YYYY-MM') = %s"

            row = db.execute(
                f"""SELECT (SELECT COUNT(*) FROM sold_orders WHERE user_id = %s AND {month_filter}) AS items,
                           (SELECT COALESCE(ROUND(SUM(item_price), 2), 0)
                            FROM sold_orders WHERE user_id = %s AND {month_filter}) AS rev,
                           COALESCE(ROUND(SUM(order_ship), 2), 0) AS ship,
                           COALESCE(ROUND(SUM(order_fees), 2), 0) AS fees
                    FROM (
                      SELECT order_id,
                             MAX(shipping) AS order_ship,
                             MAX(total_fees) AS order_fees
                      FROM sold_orders
                      WHERE user_id = %s AND {month_filter}
                      GROUP BY order_id
                    )""",
                [user_id, month_str, user_id, month_str, user_id, month_str],
            ).fetchone()
            sold_items = row["items"]
            revenue = row["rev"]
            shipping = row["ship"]
            fees = row["fees"]

            top_items = [
                dict(r) for r in db.execute(
                    f"""SELECT item_title AS title, COUNT(*) AS count,
                               ROUND(AVG(item_price), 2) AS avg_price,
                               ROUND(SUM(item_price), 2) AS revenue
                        FROM sold_orders
                        WHERE user_id = %s AND {month_filter}
                          AND item_title IS NOT NULL AND item_title != ''
                        GROUP BY item_title
                        ORDER BY count DESC LIMIT 10""",
                    [user_id, month_str],
                ).fetchall()
            ]

        active_count = 0
        if active_exists:
            active_count = db.execute(
                "SELECT COUNT(*) AS c FROM active_listings WHERE user_id = %s",
                [user_id],
            ).fetchone()["c"]

    return {
        "month": month_str,
        "available_months": available_months,
        "sold_items": sold_items,
        "revenue": revenue,
        "shipping": shipping,
        "fees": fees,
        "active_listings": active_count,
        "trends": trends,
        "top_items": top_items,
    }
